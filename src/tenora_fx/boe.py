"""UK gilt yields from the Bank of England yield-curve dataset (public Excel workbooks, no API key).

The Bank's statistical database (IADB) only carries daily par yields at 5, 10 and 20 years, so the
2-year point comes from the nominal government liability curve (GLC) workbooks the Bank publishes:

* ``latest-yield-curve-data.zip``: the current month, refreshed daily (about 200 KB);
* ``glcnominalddata.zip``: the full history back to 1979 (about 39 MB).

The archive is cached under ``data/raw/cache/boe`` together with the parsed series, and it is only
re-requested (with a conditional GET, so an unchanged file costs a single 304) when the cached
history does not cover the requested window.
"""

from __future__ import annotations

import io
import json
import logging
import re
import warnings
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from . import config, httpclient, storage
from .httpclient import HttpError

log = logging.getLogger(__name__)

YIELD_CURVE_BASE = "https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves"
LATEST_ZIP_URL = f"{YIELD_CURVE_BASE}/latest-yield-curve-data.zip"
ARCHIVE_ZIP_URL = f"{YIELD_CURVE_BASE}/glcnominalddata.zip"
ARCHIVE_NAME = "glcnominalddata.zip"
WORKBOOK_PREFIX = "GLC Nominal daily data"
SPOT_SHEET = "4. spot curve"


class BoeError(RuntimeError):
    """Bank of England yield-curve data could not be fetched or parsed."""


def _empty() -> pd.DataFrame:
    return pd.DataFrame({"value": pd.Series(dtype="float64")}, index=pd.DatetimeIndex([], name="date"))


def parse_spot_curve(xlsx: bytes, maturity: float) -> pd.DataFrame:
    """Extract one maturity (in years) from the nominal spot-curve sheet of a GLC workbook.

    Layout: a header row starting with ``years:`` lists the maturity of each column, then one row
    per date starting with a datetime. Title rows and the ``Refresh`` / ``#VALUE!`` row are
    skipped; holidays have blank cells and are dropped.
    """
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # openpyxl warns about extensions it does not support
        wb = load_workbook(io.BytesIO(xlsx), read_only=True, data_only=True)
    try:
        if SPOT_SHEET not in wb.sheetnames:
            raise BoeError(f"sheet {SPOT_SHEET!r} not found (sheets: {wb.sheetnames})")
        col: int | None = None
        rows: list[tuple[datetime, float]] = []
        for row in wb[SPOT_SHEET].iter_rows(values_only=True):
            first = row[0] if row else None
            if col is None:
                if isinstance(first, str) and first.strip().lower().startswith("years"):
                    for i, cell in enumerate(row[1:], start=1):
                        if isinstance(cell, (int, float)) and abs(float(cell) - maturity) < 1e-6:
                            col = i
                            break
                    if col is None:
                        raise BoeError(f"{maturity:g}-year column not found in {SPOT_SHEET!r}")
                continue
            if isinstance(first, datetime) and col < len(row):
                cell = row[col]
                if isinstance(cell, (int, float)):
                    rows.append((first, float(cell)))
        if col is None:
            raise BoeError(f"maturity header row not found in {SPOT_SHEET!r}")
    finally:
        wb.close()
    if not rows:
        return _empty()
    df = pd.DataFrame(rows, columns=["date", "value"]).set_index("date").sort_index()
    return df[~df.index.duplicated(keep="last")]


def workbook_years(name: str) -> tuple[int, int] | None:
    """``..._2016 to 2024.xlsx`` -> (2016, 2024); ``..._2025 to present.xlsx`` -> (2025, 9999); else None."""
    match = re.search(r"_(\d{4}) to (\d{4}|present)\.xlsx$", name)
    if not match:
        return None
    last = 9999 if match.group(2) == "present" else int(match.group(2))
    return int(match.group(1)), last


def parse_zip(src, maturity: float, years: tuple[int, int] | None = None) -> pd.DataFrame:
    """Parse every GLC Nominal workbook in a zip (path or file object), optionally only those
    whose year span overlaps ``years``."""
    frames = []
    with zipfile.ZipFile(src) as archive:
        for name in archive.namelist():
            base = Path(name).name
            if not (base.startswith(WORKBOOK_PREFIX) and base.endswith(".xlsx")):
                continue
            span = workbook_years(base)
            if years and span and (span[1] < years[0] or span[0] > years[1]):
                continue
            log.info("BoE: parsing %s", base)
            frames.append(parse_spot_curve(archive.read(name), maturity))
    if not frames:
        raise BoeError("no GLC Nominal workbook found in zip")
    df = pd.concat(frames).sort_index()
    return df[~df.index.duplicated(keep="last")]


def _download_archive(cache_dir: Path, session, timeout) -> tuple[Path, bool]:
    """Fetch the archive zip into ``cache_dir`` unless the cached copy is current. Returns (path, changed)."""
    cache_dir.mkdir(parents=True, exist_ok=True)
    zip_path = cache_dir / ARCHIVE_NAME
    meta_path = cache_dir / f"{ARCHIVE_NAME}.meta.json"
    headers = {}
    if zip_path.exists() and meta_path.exists():
        meta = json.loads(meta_path.read_text())
        if meta.get("last_modified"):
            headers["If-Modified-Since"] = meta["last_modified"]
        if meta.get("etag"):
            headers["If-None-Match"] = meta["etag"]
    resp = httpclient.get(ARCHIVE_ZIP_URL, session=session, headers=headers, timeout=timeout,
                          stream=True, label="BoE yield-curve archive")
    if resp.status_code == 304:
        log.info("BoE: archive unchanged, using cached copy")
        return zip_path, False
    tmp = zip_path.with_name(zip_path.name + ".tmp")
    with open(tmp, "wb") as fh:
        for chunk in resp.iter_content(chunk_size=1 << 20):
            fh.write(chunk)
    tmp.replace(zip_path)
    meta_path.write_text(json.dumps({
        "last_modified": resp.headers.get("Last-Modified"),
        "etag": resp.headers.get("ETag"),
        "downloaded_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
    }))
    log.info("BoE: downloaded archive (%.1f MB)", zip_path.stat().st_size / 1e6)
    return zip_path, True


def _read_json(path: Path) -> dict:
    return json.loads(path.read_text()) if path.exists() else {}


def _reaches_back(hist: pd.DataFrame | None, start_ts, parsed_from) -> bool:
    """True when the cached history covers the window start, or was parsed for a window at least
    as early (the archive itself may simply have no data that far back)."""
    if hist is None or hist.empty:
        return False
    if start_ts is None or hist.index.min() <= start_ts + pd.Timedelta(days=7):
        return True
    return parsed_from is not None and parsed_from <= start_ts


def _needs_archive(hist, latest: pd.DataFrame, start_ts, end_ts, parsed_from) -> bool:
    """True when the cached history is missing, too short, or leaves a gap before the current month."""
    if not _reaches_back(hist, start_ts, parsed_from):
        return True
    after_hist = hist.index.max() + pd.Timedelta(days=1)
    if not latest.empty:
        return len(pd.bdate_range(after_hist, latest.index.min() - pd.Timedelta(days=1))) > 0
    if end_ts is not None:
        return len(pd.bdate_range(after_hist, end_ts)) > 5
    return False


def fetch_series(
    source_id: str,
    start: date | None = None,
    end: date | None = None,
    *,
    session=None,
    cache_dir: Path | None = None,
    timeout: float = 120,
) -> pd.DataFrame:
    """Daily nominal spot yield at ``source_id`` years (e.g. ``"2.0"``) between ``start`` and ``end``."""
    try:
        maturity = float(source_id)
    except ValueError as exc:
        raise BoeError(f"{source_id!r}: expected a maturity in years, e.g. '2.0'") from exc
    cache_dir = Path(cache_dir or config.BOE_CACHE_DIR)
    start_ts = pd.Timestamp(start) if start is not None else None
    end_ts = pd.Timestamp(end) if end is not None else None

    try:
        resp = httpclient.get(LATEST_ZIP_URL, session=session, timeout=timeout, label="BoE latest yield curves")
        latest = parse_zip(io.BytesIO(resp.content), maturity)
    except (HttpError, zipfile.BadZipFile) as exc:
        raise BoeError(str(exc)) from exc

    hist_path = cache_dir / f"glc_nominal_spot_{maturity:g}y.parquet"
    hist_meta_path = hist_path.with_suffix(".meta.json")
    hist_meta = _read_json(hist_meta_path)
    parsed_from = pd.Timestamp(hist_meta["parsed_from"]) if hist_meta.get("parsed_from") else None
    hist = storage.load_parquet(hist_path) if hist_path.exists() else None
    if _needs_archive(hist, latest, start_ts, end_ts, parsed_from):
        try:
            zip_path, changed = _download_archive(cache_dir, session, timeout)
            if changed or not _reaches_back(hist, start_ts, parsed_from):
                years = (start_ts.year if start_ts is not None else 1900,
                         end_ts.year if end_ts is not None else 9999)
                parsed = parse_zip(zip_path, maturity, years=years)
                hist = parsed if hist is None else pd.concat([hist, parsed])
                window_start = start_ts if start_ts is not None else pd.Timestamp("1900-01-01")
                parsed_from = window_start if parsed_from is None else min(parsed_from, window_start)
        except (HttpError, zipfile.BadZipFile) as exc:
            raise BoeError(str(exc)) from exc

    combined = pd.concat([frame for frame in (hist, latest) if frame is not None]).sort_index()
    combined = combined[~combined.index.duplicated(keep="last")]
    combined.index.name = "date"
    storage.save_parquet(combined, hist_path)  # accumulate, so later runs rarely need the archive
    hist_meta_path.write_text(json.dumps({"parsed_from": parsed_from.strftime("%Y-%m-%d") if parsed_from is not None else None}))

    df = combined
    if start_ts is not None:
        df = df[df.index >= start_ts]
    if end_ts is not None:
        df = df[df.index <= end_ts]
    if df.empty:
        raise BoeError(f"no {maturity:g}-year observations between {start} and {end}")
    return df
